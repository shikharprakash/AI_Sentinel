"""Deck slides 1-6: title, exec summary, introduction, market, competition/Ansoff, timing."""
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.util import Inches, Pt
from deck_lib import *
import sdata as N


def s_title(prs):
    s = blank(prs, dark=True)
    text(s, 0.8, 1.7, 7.6, 0.4, "ISB PCPM CAPSTONE  |  PRODUCT PROPOSAL", size=14, color=BLUE, bold=True)
    text(s, 0.8, 2.15, 7.6, 1.6, "Product Proposal for Cisco AI Sentinel", size=40, color=WHITE, bold=True)
    text(s, 0.8, 3.85, 7.4, 1.0, "Secure, governed and smarter use of any AI tool, right inside the browser",
         size=20, color=SKY)
    text(s, 0.8, 5.55, 7, 0.9, [("Shikhar Prakash", {"bold": True, "size": 18}),
                                 ("VP of Product, Cisco  |  October 2026", {"size": 14, "color": "9FB3C8"})],
         color=WHITE)
    for (x, y, lab, col) in [(9.0, 1.7, "Govern", BLUE), (10.6, 2.6, "Protect", "0A6EA8"),
                             (9.0, 3.6, "Optimize", GREEN)]:
        c = rect(s, x, y, 2.3, 2.3, fill=col, shape=MSO_SHAPE.OVAL)
        shape_text(c, lab, size=18, color=WHITE, bold=True)
    notes(s, "Seeking approval for a 12-week MVP of Cisco AI Sentinel.")


def s_exec(prs):
    s = blank(prs)
    header(s, "Executive summary", "Make secure AI use the easiest way to use AI")
    rows = [("Problem", "Employees adopt AI tools faster than companies can govern them. 48% admit pasting "
                        "non-public data into GenAI tools."),
            ("Solution", "AI Sentinel: a browser extension + Cisco cloud that governs, protects and optimises "
                         "every prompt across ChatGPT, Gemini and more."),
            ("Why Cisco", "Duo identity, Talos intel, AI Defense models and a channel into mid-market "
                          "customers offset our late entry."),
            ("The ask", "US$1.2M for one squad and a 12-week MVP, with a paid pilot and a launch / no-launch "
                        "gate in Jan 2027.")]
    for i, (k, v) in enumerate(rows):
        y = 1.75 + i * 1.25
        badge(s, 0.6, y + 0.1, 0.6, str(i + 1), size=16)
        text(s, 1.4, y, 6.2, 0.4, k, size=16, color=NAVY, bold=True)
        text(s, 1.4, y + 0.38, 6.4, 0.8, v, size=14)
    stats = [("48%", "of employees entered non-public company data into GenAI (Cisco, 2024)"),
             ("$8.6B", "shadow-AI governance market by 2032, 35.6% CAGR (MarketsandMarkets)"),
             ("$77M", "obtainable ARR by FY2029 (SOM, author estimate)")]
    for i, (big, small) in enumerate(stats):
        y = 1.75 + i * 1.7
        card = rect(s, 8.3, y, 4.4, 1.5, fill=ICE)
        text(s, 8.5, y + 0.12, 1.9, 1.2, big, size=36, color=BLUE, bold=True, anchor="m")
        text(s, 10.35, y + 0.15, 2.25, 1.2, small, size=12, color=INK, anchor="m")


def s_intro(prs):
    s = blank(prs)
    header(s, "1  Introduction", "Cisco brings trust and reach; AI Sentinel adds the user-level AI layer")
    rect(s, 0.6, 1.7, 4.3, 5.1, fill=ICE)
    text(s, 0.85, 1.85, 3.9, 0.4, "CISCO AT A GLANCE", size=12, color=BLUE, bold=True)
    text(s, 0.85, 2.2, 3.9, 0.8, "$63.3B", size=40, color=NAVY, bold=True)
    text(s, 0.85, 2.95, 3.9, 0.4, "FY2026 revenue, +12% YoY", size=13, color=GREY)
    text(s, 0.85, 3.45, 3.9, 3.3, bullets([
        "**Networking:** Catalyst, Meraki",
        "**Security:** Secure Access, Duo, XDR, AI Defense",
        "**Observability:** Splunk, ThousandEyes",
        "**Collaboration:** Webex",
        "**Customers:** enterprises, mid-market, public sector; partner-led in 180+ countries"],
        size=14, after=6))
    text(s, 5.3, 1.7, 7.4, 1.25, [("“AI Sentinel sits quietly in the browser, catches the API key or "
                                   "customer list before it leaves, enforces company AI policy across every "
                                   "tool, and helps people write better prompts.”", {"italic": True})],
         size=16, color=NAVY)
    pillars = [("G", "Govern", "Role-based access to AI tools: allow, warn or block by user, group and app"),
               ("P", "Protect", "Real-time PII, secrets and injection checks, with one-click redaction on "
                                "the device"),
               ("O", "Optimize", "Opt-in Prompt Copilot: clearer, leaner prompts with before/after preview")]
    for i, (b, t, d) in enumerate(pillars):
        x = 5.3 + i * 2.5
        rect(s, x, 3.15, 2.3, 2.45, fill=WHITE, line=LINE)
        badge(s, x + 0.2, 3.35, 0.55, b, fill=[BLUE, "0A6EA8", GREEN][i], size=16)
        text(s, x + 0.2, 4.0, 1.95, 0.4, t, size=16, color=NAVY, bold=True)
        text(s, x + 0.2, 4.4, 1.95, 1.2, d, size=12)
    text(s, 5.3, 5.85, 7.4, 0.95, [
        ("**Buyers:** CISOs and IT leaders at 100-5,000-employee firms (Duo / Meraki base first)", {}),
        ("**Users:** knowledge workers  |  **On-ramp:** freemium for AI power users", {})], size=13)


def s_market(prs):
    s = blank(prs)
    header(s, "2  Target market landscape", "A fast-growing market with a ~$2.9B serviceable opportunity")
    yrs, vals = N.shadow_series()
    cd = CategoryChartData(); cd.categories = [str(y) for y in yrs]; cd.add_series("US$B", vals)
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.6), Inches(1.75), Inches(6.6),
                            Inches(4.6), cd)
    ch = gf.chart; chart_style(ch)
    ch.has_title = True; ch.chart_title.text_frame.text = "Shadow AI risk & governance market (US$B)"
    tp = ch.chart_title.text_frame.paragraphs[0]; tp.runs[0].font.size = Pt(14); tp.runs[0].font.bold = True
    tp.runs[0].font.color.rgb = RGB(NAVY)
    ch.has_legend = False
    pl = ch.plots[0]; pl.gap_width = 60; pl.has_data_labels = True
    dl = pl.data_labels; dl.number_format = '"$"0.0"B"'; dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END; dl.font.size = Pt(12); dl.font.color.rgb = RGB(INK)
    ser = pl.series[0]; ser.format.fill.solid(); ser.format.fill.fore_color.rgb = RGB(BLUE)
    ch.value_axis.visible = False; ch.value_axis.has_major_gridlines = False
    ch.category_axis.format.line.color.rgb = RGB(LINE)
    text(s, 0.6, 6.4, 6.6, 0.5, "35.6% CAGR 2026-2032. Source: MarketsandMarkets (Sept 2026). AI TRiSM: "
                                "$3.1B to $11.6B by 2031.", size=11, color=GREY)
    layers = [("TAM", f"${N.TAM:.0f}B", "Potential spend: ~300M enterprise GenAI seats + ~150M power users",
               SKY, NAVY, 5.5),
              ("SAM", f"${N.SAM:.1f}B", "Cisco-reachable: ~25M mid-market seats (NA, EMEA, India) + ~20M "
                                       "power users", BLUE, WHITE, 4.7),
              ("SOM", f"${N.SOM:.0f}M", "ARR by FY2029: 0.75M seats + 0.4M paying individuals", NAVY,
               WHITE, 3.9)]
    for i, (k, v, d, fill, col, w) in enumerate(layers):
        y = 1.8 + i * 1.6
        x = 7.6 + (5.5 - w) / 2
        r = rect(s, x, y, w, 1.4, fill=fill)
        text(s, x + 0.2, y + 0.1, 1.6, 0.6, k, size=14, color=col, bold=True)
        text(s, x + 0.2, y + 0.45, 1.8, 0.8, v, size=28, color=col, bold=True)
        text(s, x + 1.9, y + 0.15, w - 2.05, 1.15, d, size=12, color=col, anchor="m")
    text(s, 7.6, 6.6, 5.5, 0.35, "Bottom-up sizing: author estimates (US$72/seat/yr; US$57/user/yr)",
         size=11, color=GREY)


def s_compete(prs):
    s = blank(prs)
    header(s, "2  Competitive landscape & Ansoff", "Rivals restrict the user; nobody makes AI both safer "
                                                    "and more productive")
    x0, y0, w, h = 0.9, 1.8, 6.0, 4.6
    rect(s, x0, y0, w, h, fill=WHITE, line=LINE, shape=MSO_SHAPE.RECTANGLE)
    rect(s, x0 + w / 2, y0, w / 2, h / 2, fill=ICE, shape=MSO_SHAPE.RECTANGLE)
    ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x0 + w / 2), Inches(y0), Inches(x0 + w / 2),
                                Inches(y0 + h)); ln.line.color.rgb = RGB(LINE)
    ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x0), Inches(y0 + h / 2), Inches(x0 + w),
                                Inches(y0 + h / 2)); ln.line.color.rgb = RGB(LINE)
    pts = [("AI Sentinel", 8.4, 8.6, True), ("Prisma Browser", 7.6, 3.1, False), ("Island", 7.0, 2.0, False),
           ("Prompt Security", 6.0, 3.9, False), ("Harmonic / LayerX", 5.6, 1.4, False),
           ("MS Purview + Edge", 6.9, 5.6, False), ("Zscaler / Netskope", 7.2, 0.7, False),
           ("PromptPerfect / AIPRM", 1.0, 8.0, False), ("Native AI settings", 1.4, 4.3, False)]
    for name, px, py, hero in pts:
        cx = x0 + px / 10 * w; cy = y0 + h - py / 10 * h
        d = 0.36 if hero else 0.2
        rect(s, cx - d / 2, cy - d / 2, d, d, fill=BLUE if hero else GREY, shape=MSO_SHAPE.OVAL)
        text(s, cx + d / 2 + 0.04, cy - 0.17, 2.0, 0.34, name, size=13 if hero else 11,
             color=NAVY if hero else INK, bold=hero)
    text(s, x0, y0 + h + 0.05, w, 0.35, "Security & governance depth →", size=12, color=GREY, align="c")
    yl = text(s, x0 - 2.6, y0 + h / 2 - 0.17, 4.6, 0.34, "Productivity value to the user →", size=12,
              color=GREY, align="c")
    yl.rotation = 270
    gx, gy, cw, chh = 7.7, 2.15, 2.5, 2.0
    text(s, gx, 1.75, 5.1, 0.35, "ANSOFF MATRIX", size=12, color=BLUE, bold=True)
    text(s, gx, gy - 0.02, cw, 0.3, "Existing products", size=11, color=GREY, align="c")
    text(s, gx + cw + 0.1, gy - 0.02, cw, 0.3, "New products", size=11, color=GREY, align="c")
    cells = [("Market penetration", "Cross-sell AI Defense / Secure Access", ICE, INK),
             ("Product development", "AI Sentinel for Duo, Meraki & Secure Access customers (primary)", BLUE, WHITE),
             ("Market development", "Take Duo / Secure Access into new regions", ICE, INK),
             ("Related diversification", "AI Sentinel Personal for prosumers (contained bet)", "0A6EA8", WHITE)]
    for i, (t, d, f, c) in enumerate(cells):
        x = gx + (i % 2) * (cw + 0.1); y = gy + 0.3 + (i // 2) * (chh + 0.1)
        r = rect(s, x, y, cw, chh, fill=f)
        shape_text(r, [(t, {"bold": True, "size": 14}), (d, {"size": 12})], color=c, align="l", anchor="m")
    for k, lab in enumerate(["Existing markets", "New markets"]):
        yl = text(s, gx - 1.2, gy + 0.3 + k * (chh + 0.1) + chh / 2 - 0.16, 2.0, 0.32, lab, size=11,
                  color=GREY, align="c")
        yl.rotation = 270


def s_timing(prs):
    s = blank(prs)
    header(s, "3  Timing", "We enter third; Cisco's reach can recover most of the late-entry penalty")
    base = N.entry_shares(); adj = N.cisco_adjusted()
    cd = CategoryChartData(); cd.categories = ["Pioneer", "2nd entrant", "Cisco (3rd)"]
    cd.add_series("Order of entry only", [round(v * 100, 1) for v in base])
    cd.add_series("With Cisco preference 1.2x, marketing 1.5x", [round(v * 100, 1) for v in adj[:3]])
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.6), Inches(1.75), Inches(6.6),
                            Inches(4.7), cd)
    ch = gf.chart; chart_style(ch)
    ch.has_title = True; ch.chart_title.text_frame.text = "Long-run market share, % (Urban et al., 1986)"
    r0 = ch.chart_title.text_frame.paragraphs[0].runs[0]; r0.font.size = Pt(14); r0.font.bold = True
    r0.font.color.rgb = RGB(NAVY)
    ch.has_legend = True; ch.legend.position = XL_LEGEND_POSITION.BOTTOM; ch.legend.include_in_layout = False
    ch.legend.font.size = Pt(12)
    pl = ch.plots[0]; pl.gap_width = 70; pl.overlap = -10; pl.has_data_labels = True
    pl.data_labels.number_format = '0.0"%"'; pl.data_labels.number_format_is_linked = False
    pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END; pl.data_labels.font.size = Pt(12)
    for ser, col in zip(pl.series, [GREY, BLUE]):
        ser.format.fill.solid(); ser.format.fill.fore_color.rgb = RGB(col)
    ch.value_axis.visible = False; ch.value_axis.has_major_gridlines = False
    ch.category_axis.format.line.color.rgb = RGB(LINE)
    gap = (base[0] - base[2]); seg = N.CONTESTED_SEGMENT_M
    r = rect(s, 7.6, 1.8, 2.45, 1.6, fill=ICE)
    shape_text(r, [(f"-{gap * 100:.0f} pts", {"size": 30, "bold": True, "color": "C0392B"}),
                   (f"share vs. pioneer = ~${gap * seg:.0f}M ARR lost by FY2029", {"size": 12})])
    r = rect(s, 10.25, 1.8, 2.45, 1.6, fill=ICE)
    shape_text(r, [(f"+{(adj[2] - base[2]) * 100:.1f} pts", {"size": 30, "bold": True, "color": BLUE}),
                   (f"recovered: {adj[2]:.0%} share, ~${adj[2] * seg:.0f}M ARR", {"size": 12})])
    text(s, 7.6, 3.65, 5.1, 0.4, "Late-mover advantages we will use", size=16, color=NAVY, bold=True)
    text(s, 7.6, 4.1, 5.1, 2.6, bullets([
        "Pioneers already educated the buyer; budgets exist",
        "Learn from rivals: heavy rollouts and noisy alerts",
        "Duo / Meraki base and partner channel cut CAC",
        "Bundling and brand trust in Enterprise Agreements",
        "Window closing as SSE vendors bundle: ship in 12 weeks"], size=14, after=5))
    text(s, 0.6, 6.5, 6.6, 0.4, f"Contested segment: browser-layer AI security, mid-market, ~${seg}M/yr by "
                                "FY2029 (estimate)", size=11, color=GREY)


SLIDES = [s_title, s_exec, s_intro, s_market, s_compete, s_timing]
