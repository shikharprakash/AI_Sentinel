"""python-pptx helpers for the AI Sentinel executive deck (built on the ISB template's master)."""
import copy
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

RGB = lambda h: RGBColor.from_string(h)
NAVY, BLUE, SKY, ICE = "0D274D", "049FD9", "BFE6F7", "EEF7FC"
GREEN, AMBER, GREY, INK, WHITE, LINE = "6CC04A", "FBAB18", "5B6770", "1F2A36", "FFFFFF", "D5DDE3"
FONT = "Calibri"
W, H = 13.333, 7.5
STATE = {"n": 0}


def blank(prs, dark=False):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    STATE["n"] += 1
    bg = s.background.fill
    bg.solid(); bg.fore_color.rgb = RGB(NAVY if dark else WHITE)
    return s


def text(slide, x, y, w, h, paras, size=14, color=INK, bold=False, align="l", anchor="t",
         font=FONT, spacing=4, margin=0.0, name=None):
    """paras: str | list of str | list of (str|[(run, {opts})], {para opts})."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    m = Inches(margin)
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = m
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    fill_tf(tf, paras, size, color, bold, align, font, spacing)
    return tb


def _bullet(p, char="•", color=BLUE, indent=0.22):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(indent)))); pPr.set("indent", str(-int(Inches(indent))))
    for tag in ("a:buClr", "a:buChar", "a:buNone"):
        for e in pPr.findall(qn(tag)):
            pPr.remove(e)
    clr = etree.SubElement(pPr, qn("a:buClr")); sc = etree.SubElement(clr, qn("a:srgbClr")); sc.set("val", color)
    bu = etree.SubElement(pPr, qn("a:buChar")); bu.set("char", char)


def fill_tf(tf, paras, size=14, color=INK, bold=False, align="l", font=FONT, spacing=4):
    if isinstance(paras, str):
        paras = [paras]
    first = True
    for item in paras:
        opts = {}
        if isinstance(item, tuple):
            item, opts = item
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[opts.get("align", align)]
        p.space_after = Pt(opts.get("after", spacing))
        if opts.get("bullet"):
            _bullet(p, color=opts.get("bucolor", BLUE), indent=opts.get("indent", 0.22))
        runs = item if isinstance(item, list) else _parse(item)
        for rtext, ropts in runs:
            r = p.add_run(); r.text = rtext
            f = r.font
            f.name = font; f.size = Pt(ropts.get("size", opts.get("size", size)))
            f.bold = ropts.get("bold", opts.get("bold", bold))
            f.italic = ropts.get("italic", opts.get("italic", False))
            f.color.rgb = RGB(ropts.get("color", opts.get("color", color)))


def _parse(s):
    """'**bold**' inline markup -> runs."""
    out = []
    for i, chunk in enumerate(s.split("**")):
        if chunk:
            out.append((chunk, {"bold": True} if i % 2 else {}))
    return out


def bullets(items, size=14, color=INK, after=6):
    return [(it, {"bullet": True, "size": size, "color": color, "after": after}) for it in items]


def rect(slide, x, y, w, h, fill=ICE, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06, name=None):
    sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        sh.name = name
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = RGB(fill)
    if line:
        sh.line.color.rgb = RGB(line); sh.line.width = Pt(1)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = radius
    sh.text_frame.margin_left = sh.text_frame.margin_right = Inches(0.08)
    return sh


def shape_text(sh, paras, size=14, color=INK, bold=False, align="c", anchor="m"):
    tf = sh.text_frame; tf.word_wrap = True
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    fill_tf(tf, paras, size, color, bold, align)
    return sh


def badge(slide, x, y, d, label, fill=BLUE, color=WHITE, size=16):
    c = rect(slide, x, y, d, d, fill=fill, shape=MSO_SHAPE.OVAL)
    c.text_frame.margin_left = c.text_frame.margin_right = 0
    shape_text(c, label, size=size, color=color, bold=True)
    return c


def header(slide, kicker, title, dark=False):
    text(slide, 0.6, 0.38, 12.1, 0.35, kicker.upper(), size=12, color=BLUE, bold=True, name="Kicker")
    text(slide, 0.6, 0.68, 12.1, 0.8, title, size=28, color=WHITE if dark else NAVY, bold=True,
         name="Title")
    footer(slide, dark)


def footer(slide, dark=False):
    col = "9FB3C8" if dark else GREY
    text(slide, 0.6, 7.05, 8, 0.3, "Cisco AI Sentinel  |  Product proposal  |  Shikhar Prakash", size=10,
         color=col, name="Footer")
    text(slide, 11.73, 7.05, 1.0, 0.3, str(STATE["n"]), size=10, color=col, align="r", name="Page")


def table(slide, x, y, w, rows, col_w, size=12, row_h=0.36, header_fill=NAVY, first_bold=True, zebra=True):
    nr, nc = len(rows), len(rows[0])
    gt = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(row_h * nr))
    tbl = gt.table
    tblPr = tbl._tbl.tblPr
    style = tblPr.find(qn("a:tableStyleId"))
    if style is not None:
        style.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"  # No style, table grid
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Inches(cw)
    for i in range(nr):
        tbl.rows[i].height = Inches(row_h)
        for j in range(nc):
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            fill = header_fill if i == 0 else (ICE if zebra and i % 2 == 0 else WHITE)
            cell.fill.solid(); cell.fill.fore_color.rgb = RGB(fill)
            tf = cell.text_frame; tf.word_wrap = True
            val = rows[i][j]
            fill_tf(tf, val if isinstance(val, list) else str(val), size=size,
                    color=WHITE if i == 0 else INK, bold=(i == 0) or (first_bold and j == 0))
            _cell_border(cell)
    return gt


def _cell_border(cell, color=LINE, w=6350):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        ln = etree.SubElement(tcPr, qn(tag)); ln.set("w", str(w))
        sf = etree.SubElement(ln, qn("a:solidFill")); c = etree.SubElement(sf, qn("a:srgbClr")); c.set("val", color)
    # schema order: borders must precede fill
    fills = [e for e in tcPr if e.tag in (qn("a:solidFill"), qn("a:noFill"))]
    for f in fills:
        tcPr.remove(f); tcPr.append(f)


def chart_style(chart, size=12):
    chart.font.size = Pt(size); chart.font.name = FONT
    chart.font.color.rgb = RGB(GREY)


def notes(slide, s):
    slide.notes_slide.notes_text_frame.text = s
