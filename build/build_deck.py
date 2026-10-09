"""Build the executive deck on the ISB PowerPoint template (keeps its master/theme, replaces slides)."""
import os
import sys
from pptx import Presentation
import deck_lib as D
import deck_a
try:
    import deck_b
except ImportError:
    deck_b = None
try:
    import deck_c
except ImportError:
    deck_c = None

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "ISB-PCPM-CapstonePowerpointTemplate.pptx")
DST = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, "AI_Sentinel_Capstone_Deck.pptx")

prs = Presentation(SRC)
# drop the template's placeholder slides (incl. the "Notes - delete" slide)
sldIdLst = prs.slides._sldIdLst
for sldId in list(sldIdLst):
    prs.part.drop_rel(sldId.rId)
    sldIdLst.remove(sldId)

for mod in (deck_a, deck_b, deck_c):
    if mod:
        for fn in mod.SLIDES:
            fn(prs)

prs.core_properties.author = "Shikhar Prakash"
prs.core_properties.title = "Product Proposal for Cisco AI Sentinel"
prs.save(DST)
print("saved", DST, len(prs.slides), "slides")
