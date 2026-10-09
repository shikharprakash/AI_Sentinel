"""Deck slides 7-12: customers (2), product (2), concept testing (2)."""
import os
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt
from deck_lib import *
import sdata as N

IMG = os.path.join(os.path.dirname(__file__), "img")


def s_personas(prs):
    s = blank(prs)
    header(s, "4  Customers", "Three personas, one journey from 'unaware' to advocate")
    ps = [("A", "Ananya, 27", "AI power user | Developer, Bengaluru",
           "**Wants:** precise answers, fewer tokens\n**Pain:** vague prompts; once pasted an API key\n"
           "**Role:** user and buyer (Rs 299-599/mo)"),
          ("B", "Vikram, 44", "Security leader | CISO, 1,200-staff IT firm",
           "**Wants:** safe AI adoption, audit evidence\n**Pain:** no visibility; bans backfire; noisy DLP\n"
           "**Role:** economic buyer"),
          ("C", "Meera, 33", "Knowledge worker | HR partner, fintech",
           "**Wants:** finish writing faster, never leak PII\n**Pain:** unsure what is 'sensitive'; cryptic "
           "blocks\n**Role:** user and influencer")]
    for i, (b, name, role, body) in enumerate(ps):
        x = 0.6 + i * 4.1
        rect(s, x, 1.7, 3.9, 2.75, fill=ICE)
        badge(s, x + 0.2, 1.88, 0.6, b, fill=[GREEN, NAVY, BLUE][i], size=18)
        text(s, x + 0.95, 1.85, 2.85, 0.4, name, size=16, color=NAVY, bold=True)
        text(s, x + 0.95, 2.22, 2.85, 0.5, role, size=12, color=GREY)
        text(s, x + 0.2, 2.8, 3.55, 1.6, [(ln, {}) for ln in body.split("\n")], size=13, spacing=5)
    stages = [("Unaware", "AI-leak news, Talos blogs"), ("Aware", "Account team, partner, Web Store"),
              ("Consider", "Demo, free tier, reviews"), ("Pilot", "4 weeks: monitor, then enforce"),
              ("Purchase", "Partner / EA / in-app"), ("Use", "MDM deploy, inline tips"),
              ("Advocate", "Case study, referrals")]
    text(s, 0.6, 4.7, 12, 0.35, "CUSTOMER JOURNEY (ENTERPRISE; INDIVIDUALS SKIP THE PILOT)", size=12,
         color=BLUE, bold=True)
    for i, (st, tp) in enumerate(stages):
        x = 0.6 + i * 1.75
        ch = rect(s, x, 5.1, 1.85, 0.62, fill=[SKY, SKY, BLUE, BLUE, NAVY, NAVY, GREEN][i],
                  shape=MSO_SHAPE.CHEVRON if i else MSO_SHAPE.PENTAGON)
        shape_text(ch, st, size=13, bold=True, color=NAVY if i < 2 else WHITE)
        text(s, x + 0.1, 5.82, 1.6, 0.9, tp, size=11, color=INK)


def s_research(prs):
    s = blank(prs)
    header(s, "4  Customer research", "Protection without friction wins; buyers and users value different "
                                       "things")
    rows = sorted(N.rating_table(), key=lambda d: d["mean"])
    cd = CategoryChartData(); cd.categories = [d["short"] for d in rows]
    cd.add_series("Mean importance", [d["mean"] for d in rows])
    gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.6), Inches(1.7), Inches(6.4),
                            Inches(4.85), cd)
    ch = gf.chart; chart_style(ch); ch.has_legend = False
    ch.has_title = True; ch.chart_title.text_frame.text = "Importance rating, mean of 5 interviews (1-5)"
    r0 = ch.chart_title.text_frame.paragraphs[0].runs[0]; r0.font.size = Pt(14); r0.font.bold = True
    r0.font.color.rgb = RGB(NAVY)
    pl = ch.plots[0]; pl.gap_width = 45; pl.has_data_labels = True
    pl.data_labels.number_format = "0.0"; pl.data_labels.number_format_is_linked = False
    pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END; pl.data_labels.font.size = Pt(12)
    ser = pl.series[0]
    for i, d in enumerate(rows):
        pt = ser.points[i]; pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = RGB(NAVY if d["mean"] >= 4.2 else BLUE)
    ch.value_axis.visible = False; ch.value_axis.has_major_gridlines = False
    ch.value_axis.maximum_scale = 5.4; ch.value_axis.minimum_scale = 0
    ch.category_axis.format.line.color.rgb = RGB(LINE); ch.category_axis.tick_labels.font.size = Pt(12)
    text(s, 0.6, 6.6, 6.4, 0.35, "Illustrative data: replace with real interview results", size=11,
         color=GREY)
    key = {d["short"]: d for d in N.rating_table()}
    trows = [["Feature", "Enterprise", "Individual"]]
    for f in ["Tool access (RBAC)", "Audit & reports", "Prompt optimizer", "Preview & undo"]:
        trows.append([f, f"{key[f]['ent']:.1f}", f"{key[f]['ind']:.1f}"])
    text(s, 7.4, 1.7, 5.3, 0.35, "SEGMENT SPLIT (MEAN SCORE)", size=12, color=BLUE, bold=True)
    table(s, 7.4, 2.1, 5.3, trows, [2.7, 1.3, 1.3], size=12, row_h=0.38)
    text(s, 7.4, 4.2, 5.3, 2.6, bullets([
        "**Top 2 (4.6):** data redaction and low friction. Latency is a must-have",
        "**Local-first (4.4)** matters to everyone: lead with 'checked on your device'",
        "**Split** confirms one core, two packages: governance for buyers, Copilot for users",
        "**Savings dashboard (3.4)** moves to P1"], size=13, after=6))


def s_product(prs):
    s = blank(prs)
    header(s, "5  Product", "Highest opportunity at moderate risk: build it, de-risked through an MVP")
    rows = [["Criterion (1-5)", "A. AI Sentinel", "B. Secure Access AI-DLP", "C. Prompt optimiser"],
            ["Market opportunity", "5", "4", "3"], ["Strategic fit", "4", "5", "2"],
            ["Differentiation", "5", "2", "2"], ["Revenue (3 yr)", "4", "3", "2"],
            ["Technical risk (5 = low)", "3", "4", "4"], ["Market risk (5 = low)", "3", "4", "2"],
            [[("Weighted score", {"bold": True})], [("3.9", {"bold": True, "color": BLUE})], "3.7", "2.6"]]
    text(s, 0.6, 1.7, 6.6, 0.35, "OPPORTUNITY COST & DEVELOPMENT RISK (same 12-week budget)", size=12,
         color=BLUE, bold=True)
    table(s, 0.6, 2.1, 6.6, rows, [2.4, 1.3, 1.6, 1.3], size=12, row_h=0.42)
    text(s, 0.6, 5.65, 6.6, 1.2, "Weighted: opportunity average x 60% + risk average x 40%. B stays in the "
                                 "plan as a later Secure Access integration.", size=12, color=GREY)
    text(s, 7.6, 1.7, 5.1, 0.35, "PRODUCT / COMPANY FIT", size=12, color=BLUE, bold=True)
    fits = [("+", "Duo identity, Talos intel, AI Defense models, Splunk, channel", GREEN),
            ("–", "Gaps: PLG / consumer motion, extension UX, prompt-quality ML", AMBER)]
    for i, (b, t, c) in enumerate(fits):
        y = 2.1 + i * 0.75
        badge(s, 7.6, y + 0.05, 0.45, b, fill=c, size=16)
        text(s, 8.2, y, 4.5, 0.7, t, size=13, anchor="m")
    text(s, 7.6, 3.75, 5.1, 0.35, "PRODUCT / MARKET FIT: HYPOTHESES TO TEST", size=12, color=BLUE, bold=True)
    text(s, 7.6, 4.15, 5.1, 2.7, bullets([
        "Users prefer 'Redact & send' over a hard block",
        "Friction acceptable: p95 under 100 ms, one click",
        "40%+ accept optimiser suggestions",
        "Mid-market pays US$6-8/user/month without SSE",
        "A browser extension is an acceptable MVP control"], size=13, after=5))


def s_product2(prs):
    s = blank(prs)
    header(s, "5  Product strategy", "Positioned as security employees want installed")
    ps = [("Product", "Extension + console; Govern, Protect, Optimize; 4 tiers"),
          ("Price", "Freemium + per-seat + metered advanced ops"),
          ("Place", "Chrome / Edge stores; Cisco account teams, partners, marketplaces, EA"),
          ("Promotion", "Talos AI-leak research, Cisco Live, partner enablement, communities")]
    text(s, 0.6, 1.7, 6, 0.35, "MARKETING MIX (4Ps)", size=12, color=BLUE, bold=True)
    for i, (t, d) in enumerate(ps):
        x = 0.6 + (i % 2) * 3.05; y = 2.1 + (i // 2) * 1.55
        r = rect(s, x, y, 2.9, 1.4, fill=ICE)
        shape_text(r, [(t, {"bold": True, "size": 15, "color": NAVY}), (d, {"size": 12})], align="l")
    r = rect(s, 0.6, 5.35, 5.95, 1.5, fill=NAVY)
    shape_text(r, [("POSITIONING", {"bold": True, "size": 11, "color": BLUE}),
                   ("For AI-using organisations, AI Sentinel is the browser-native layer that unites "
                    "governance, data protection and prompt optimisation, so the secure way is the most "
                    "productive way.", {"size": 13, "color": WHITE, "italic": True})], align="l")
    text(s, 7.0, 1.7, 5.7, 0.35, "THREE LEVELS OF THE PRODUCT", size=12, color=BLUE, bold=True)
    for d, col in [(3.1, SKY), (2.15, BLUE), (1.2, NAVY)]:
        off = (3.1 - d) / 2
        rect(s, 7.0 + off, 2.1 + off, d, d, fill=col, shape=MSO_SHAPE.OVAL)
    lv = [("Core", "Confidence to use any AI tool safely, with better results", NAVY),
          ("Actual", "Extension, inline redaction, Copilot, console, Cisco brand", BLUE),
          ("Augmented", "Duo SSO, Splunk/XDR, Talos packs, audit reports, support", "4BA3C7")]
    for i, (t, d, c) in enumerate(lv):
        y = 2.15 + i * 1.0
        text(s, 10.3, y, 2.4, 0.95, [(t, {"bold": True, "color": c, "size": 14}), (d, {"size": 12})])
    text(s, 7.0, 5.35, 5.7, 0.35, "BUSINESS & REVENUE MODEL", size=12, color=BLUE, bold=True)
    text(s, 7.0, 5.7, 5.7, 1.2, bullets([
        "B2B SaaS with a product-led B2C on-ramp",
        "Seats + metered ops + Personal Plus + add-ons; 75%+ GM target"], size=13, after=4))


def s_concept(prs):
    s = blank(prs)
    header(s, "6  Concept testing", "Lo-fi prototype tested 1-on-1 with three target customers")
    shots = [("wf_protect.png", "Protect: inline detection + 'Redact & send'"),
             ("wf_optimize.png", "Optimize: Copilot before/after, accept / revert"),
             ("wf_admin.png", "Govern: group policies, monitor / enforce, events")]
    for i, (f, cap) in enumerate(shots):
        x = 0.6 + i * 4.1
        pic = s.shapes.add_picture(os.path.join(IMG, f), Inches(x), Inches(1.75), width=Inches(3.9))
        text(s, x, 1.85 + pic.height / 914400, 3.9, 0.5, cap, size=12, color=NAVY, bold=True, align="c")
    rect(s, 0.6, 4.75, 12.1, 2.05, fill=ICE)
    text(s, 0.85, 4.9, 3.6, 0.35, "METHOD", size=12, color=BLUE, bold=True)
    text(s, 0.85, 5.25, 3.7, 1.5, "Qualitative think-aloud sessions (45 min) + 10-question survey with "
                                  "Likert, Sean Ellis and Van Westendorp questions", size=13)
    text(s, 4.8, 4.9, 3.6, 0.35, "PARTICIPANTS", size=12, color=BLUE, bold=True)
    text(s, 4.8, 5.25, 3.7, 1.5, bullets(["Ananya: developer (Persona A)", "Vikram: CISO (Persona B)",
                                          "Meera: HR partner (Persona C)"], size=13, after=3))
    text(s, 8.75, 4.9, 3.8, 0.35, "TASKS", size=12, color=BLUE, bold=True)
    text(s, 8.75, 5.25, 3.8, 1.5, bullets(["Send a prompt containing a PAN + API key",
                                           "Improve a vague prompt", "Create an HR policy (CISO)"],
                                          size=13, after=3))


def s_results(prs):
    s = blank(prs)
    header(s, "6  Concept test results", "Validated: clear, trusted and wanted, if it fixes rather than "
                                          "blocks")
    stats = [("3/3", "completed the redaction task unaided"), ("4.7/5", "warning clarity"),
             ("67%", "'very disappointed' without it (PMF bar: 40%)"), ("Rs 399", "Personal price; "
                                                                        "US$6-8/user for business")]
    for i, (b, t) in enumerate(stats):
        x = 0.6 + i * 3.05
        r = rect(s, x, 1.7, 2.9, 1.45, fill=ICE)
        text(s, x + 0.15, 1.78, 2.6, 0.7, b, size=30, color=BLUE, bold=True)
        text(s, x + 0.15, 2.45, 2.65, 0.65, t, size=12)
    rows = [["What we learned", "What we change"],
            ["Meera hesitated at 'Send anyway'", "'Redact & send' is the default; overrides need a reason"],
            ["Auto pop-ups annoyed Ananya", "Copilot on demand only; add a 'Concise' mode"],
            ["Vikram won't block on day one", "Monitor mode default for pilots; one-click enforce"],
            ["Auditors need evidence", "Metadata-only event export (CSV / Splunk) in the MVP"],
            ["On-device message missed (3.7/5)", "Persistent 'Checked on device' badge"]]
    table(s, 0.6, 3.45, 12.1, rows, [5.0, 7.1], size=13, row_h=0.5)
    text(s, 0.6, 6.55, 12.1, 0.35, "Illustrative results (n = 3): replace with real sessions before "
                                   "submission", size=11, color=GREY)


SLIDES = [s_personas, s_research, s_product, s_product2, s_concept, s_results]
