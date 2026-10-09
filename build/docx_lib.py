"""Helpers to render content blocks into a bordered single-cell 'answer box' in the Word template."""
import os
import re
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

IMG = os.path.join(os.path.dirname(__file__), "img")
NAVY = RGBColor(0x0D, 0x27, 0x4D)
BLUE = RGBColor(0x04, 0x9F, 0xD9)
GREY = RGBColor(0x5B, 0x67, 0x70)
BODY_PT = 10.5
BOX_W = 6.2  # inches available inside A4 margins


def _border_el(tag, color="049FD9", sz="8"):
    el = OxmlElement(f"w:{tag}")
    el.set(qn("w:val"), "single"); el.set(qn("w:sz"), sz)
    el.set(qn("w:space"), "0"); el.set(qn("w:color"), color)
    return el


def set_table_borders(tbl, color="049FD9", sz="8", inner=False):
    tblPr = tbl._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for tag in ["top", "left", "bottom", "right"] + (["insideH", "insideV"] if inner else []):
        borders.append(_border_el(tag, color, sz))
    tblPr.append(borders)


def set_cell_shading(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{k}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); mar.append(e)
    tcPr.append(mar)


def set_width(cell, inches):
    cell.width = Inches(inches)


def fix_grid(tbl, widths):
    """Fixed layout with explicit grid so Word/LibreOffice honour column widths."""
    tbl.autofit = False
    t = tbl._tbl
    tblPr = t.tblPr
    for old in tblPr.findall(qn("w:tblW")) + tblPr.findall(qn("w:tblLayout")):
        tblPr.remove(old)
    tw = OxmlElement("w:tblW"); tw.set(qn("w:w"), str(int(sum(widths) * 1440))); tw.set(qn("w:type"), "dxa")
    tblPr.append(tw)
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    grid = t.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(w * 1440)))


def add_runs(par, text, size=BODY_PT, color=None, italic=False, bold_all=False):
    """Supports **bold** inline markup."""
    for i, chunk in enumerate(re.split(r"\*\*", text)):
        if not chunk:
            continue
        r = par.add_run(chunk)
        r.font.size = Pt(size)
        r.bold = bold_all or (i % 2 == 1)
        r.italic = italic
        if color is not None:
            r.font.color.rgb = color
    return par


def fmt(par, before=0, after=4, line=1.08, indent=None, hanging=None, align=None):
    pf = par.paragraph_format
    pf.space_before = Pt(before); pf.space_after = Pt(after); pf.line_spacing = line
    if indent is not None:
        pf.left_indent = Inches(indent)
    if hanging is not None:
        pf.first_line_indent = Inches(-hanging)
    if align == "center":
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return par


class Box:
    """Renders blocks into the cell of a one-cell table."""

    def __init__(self, cell):
        self.cell = cell
        self.first = True

    def _p(self):
        if self.first:
            self.first = False
            return self.cell.paragraphs[0]
        return self.cell.add_paragraph()

    def h(self, text):
        was_first = self.first
        p = fmt(self._p(), before=8, after=3)
        p.paragraph_format.keep_with_next = not was_first
        add_runs(p, text, size=12, color=NAVY, bold_all=True)

    def h2(self, text):
        p = fmt(self._p(), before=4, after=2)
        p.paragraph_format.keep_with_next = True
        add_runs(p, text, size=BODY_PT + 0.5, color=BLUE, bold_all=True)

    def p(self, text):
        add_runs(fmt(self._p(), after=5), text)

    def note(self, text):
        p = fmt(self._p(), before=2, after=6)
        add_runs(p, text, size=9, color=GREY, italic=True)

    def b(self, items, level=0):
        for it in items:
            if isinstance(it, (list, tuple)):
                self.b(it, level + 1)
                continue
            p = fmt(self._p(), after=2, indent=0.25 + 0.25 * level, hanging=0.18)
            add_runs(p, ("•" if level == 0 else "–") + "\t" + it)
            p.paragraph_format.tab_stops.add_tab_stop(Inches(0.25 + 0.25 * level))

    def n(self, items):
        for i, it in enumerate(items, 1):
            p = fmt(self._p(), after=2, indent=0.3, hanging=0.22)
            add_runs(p, f"{i}.\t" + it)
            p.paragraph_format.tab_stops.add_tab_stop(Inches(0.3))

    def img(self, name, width=5.8, caption=None):
        p = fmt(self._p(), before=4, after=2, align="center")
        p.add_run().add_picture(os.path.join(IMG, name), width=Inches(width))
        if caption:
            add_runs(fmt(self._p(), after=6, align="center"), caption, size=8.5, color=GREY, italic=True)

    def t(self, header, rows, widths=None, size=9, first_bold=True):
        if not self.first:
            fmt(self.cell.add_paragraph(), after=0)
        self.first = False
        tbl = self.cell.add_table(rows=1 + len(rows), cols=len(header))
        set_table_borders(tbl, color="B7C3CC", sz="4", inner=True)
        widths = widths or [BOX_W / len(header) - 0.05] * len(header)
        scale = min(1.0, (BOX_W - 0.15) / sum(widths))
        widths = [w * scale for w in widths]
        fix_grid(tbl, widths)
        for r_i, row in enumerate([header] + rows):
            for c_i, val in enumerate(row):
                cell = tbl.cell(r_i, c_i)
                set_width(cell, widths[c_i]); set_cell_margins(cell, 40, 40, 70, 70)
                par = fmt(cell.paragraphs[0], after=0, line=1.0)
                if r_i == 0:
                    set_cell_shading(cell, "0D274D")
                    add_runs(par, str(val), size=size, color=RGBColor(255, 255, 255), bold_all=True)
                else:
                    if r_i % 2 == 0:
                        set_cell_shading(cell, "EEF7FC")
                    for k, line in enumerate(str(val).split("\n")):
                        lp = par if k == 0 else fmt(cell.add_paragraph(), after=0, line=1.0)
                        add_runs(lp, line, size=size, bold_all=(first_bold and c_i == 0))
        nrows = len(rows) + 1
        for r_i, row in enumerate(tbl.rows):
            trPr = row._tr.get_or_add_trPr()
            trPr.append(OxmlElement("w:cantSplit"))
            if r_i == 0:
                trPr.append(OxmlElement("w:tblHeader"))
        fmt(self.cell.add_paragraph(), after=2)
        self.first = False

    def render(self, blocks):
        for blk in blocks:
            kind, *args = blk
            getattr(self, kind)(*args)


TBLPR_ORDER = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize",
               "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar",
               "tblLook", "tblCaption", "tblDescription"]
TCPR_ORDER = ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar",
              "textDirection", "tcFitText", "vAlign", "hideMark"]


def normalize(body):
    """Sort tblPr / tcPr children into schema order (python-docx appends in call order)."""
    for tag, order in (("w:tblPr", TBLPR_ORDER), ("w:tcPr", TCPR_ORDER)):
        rank = {qn("w:" + n): i for i, n in enumerate(order)}
        for el in body.iter(qn(tag)):
            kids = list(el)
            kids.sort(key=lambda k: rank.get(k.tag, 99))
            for k in kids:
                el.remove(k)
            for k in kids:
                el.append(k)
