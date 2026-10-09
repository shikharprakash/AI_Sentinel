"""Fill the ISB capstone Word template: replace the 10 fixed text boxes with flowing answer boxes."""
import os
import re
import sys
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

import docx_lib as L
from content_a import S1, S2, S3
from content_b import S4, S5
from content_c import S6, S7, S8
from content_d import S9, S10, SOURCES

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "CapstoneWordTemplate (2).docx")
DST = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, "AI_Sentinel_Capstone_Report.docx")
SECTIONS = [S1, S2, S3, S4, S5, S6, S7, S8, S9, S10]
MC = "{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent"

doc = docx.Document(SRC)
body = doc.element.body

# 1. Remove every floating text box (mc:AlternateContent / drawing / pict) and emptied runs.
for el in list(body.iter(MC, qn("w:drawing"), qn("w:pict"))):
    parent = el.getparent()
    if parent is None:
        continue
    parent.remove(el)
    if parent.tag == qn("w:r") and not parent.xpath("./w:t|./w:tab|./w:br"):
        parent.getparent().remove(parent)
assert not list(body.iter(MC)), "text boxes remain"


def text_of(p_el):
    return "".join(t.text or "" for t in p_el.iter(qn("w:t"))).strip()


# 2. Locate section headings.
paras = [el for el in body if el.tag == qn("w:p")]
heads = [p for p in paras if re.match(r"Section \d+:", text_of(p))]
assert len(heads) == 10, len(heads)


def new_box():
    tbl = doc.add_table(rows=1, cols=1)
    L.set_table_borders(tbl, color="049FD9", sz="8")
    cell = tbl.cell(0, 0)
    L.fix_grid(tbl, [L.BOX_W + 0.07])
    L.set_width(cell, L.BOX_W + 0.07)
    L.set_cell_margins(cell, 120, 120, 160, 160)
    L.set_cell_shading(cell, "FFFFFF")
    return tbl


# 3. For each section: drop empty spacer paragraphs, insert answer box after the instructions.
children = list(body)
for i, head in enumerate(heads):
    start = children.index(head)
    end = children.index(heads[i + 1]) if i + 1 < len(heads) else len(children) - 1  # sectPr last
    block = [el for el in children[start + 1:end] if el.tag == qn("w:p")]
    instr = [p for p in block if text_of(p)]
    for p in block:
        if not text_of(p):
            body.remove(p)
    anchor = instr[-1] if instr else head
    tbl = new_box()
    anchor.addnext(tbl._tbl)
    L.Box(tbl.cell(0, 0)).render(SECTIONS[i])
    spacer = OxmlElement("w:p")
    tbl._tbl.addnext(spacer)
    if i > 0:
        Paragraph(head, doc).paragraph_format.page_break_before = True

# 4. Title line at the top and sources at the end.
first = Paragraph(heads[0], doc)
title = first.insert_paragraph_before()
L.add_runs(title, "Product Proposal: Cisco AI Sentinel", size=18, color=L.NAVY, bold_all=True)
sub = first.insert_paragraph_before()
L.add_runs(sub, "Shikhar Prakash, VP of Product, Cisco  |  ISB PCPM Capstone  |  October 2026",
           size=10.5, color=L.GREY)
sub.paragraph_format.space_after = Pt(12)

src_head = doc.add_paragraph()
src_head.paragraph_format.page_break_before = True
L.add_runs(src_head, "Sources and notes", size=12, bold_all=True, color=L.NAVY)
for k, s in enumerate(SOURCES, 1):
    p = doc.add_paragraph()
    L.fmt(p, after=3, indent=0.3, hanging=0.3)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(0.3))
    L.add_runs(p, f"{k}.\t{s}", size=9.5)

L.normalize(body)
doc.core_properties.author = "Shikhar Prakash"
doc.core_properties.title = "Product Proposal: Cisco AI Sentinel"
doc.save(DST)
print("saved", DST)
